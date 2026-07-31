from __future__ import annotations

import shutil
import subprocess
import tempfile
from pathlib import Path

from docx import Document

EXTENSOES_PERMITIDAS = {".doc", ".docx"}
MAX_UPLOAD_BYTES = 10 * 1024 * 1024


def _extrair_docx(caminho: Path) -> str:
    documento = Document(caminho)
    partes: list[str] = []

    for paragrafo in documento.paragraphs:
        texto = paragrafo.text.strip()
        if texto:
            partes.append(texto)

    for tabela in documento.tables:
        for linha in tabela.rows:
            celulas = [celula.text.strip() for celula in linha.cells if celula.text.strip()]
            if celulas:
                partes.append(" | ".join(celulas))

    return "\n".join(partes).strip()


def _extrair_doc(caminho: Path) -> str:
    executavel = shutil.which("antiword")
    if not executavel:
        raise ValueError(
            "Arquivos .doc exigem o programa antiword instalado no servidor. "
            "Converta o contrato para .docx ou instale essa dependência."
        )

    resultado = subprocess.run(
        [executavel, str(caminho)],
        capture_output=True,
        text=True,
        timeout=30,
        check=False,
    )

    if resultado.returncode != 0:
        raise ValueError("Não foi possível extrair o texto do arquivo .doc enviado.")

    return resultado.stdout.strip()


def extract_text_from_doc(file) -> str:
    nome_original = Path(file.filename or "").name
    extensao = Path(nome_original).suffix.lower()

    if extensao not in EXTENSOES_PERMITIDAS:
        raise ValueError("Envie somente um contrato nos formatos .doc ou .docx.")

    conteudo = file.file.read(MAX_UPLOAD_BYTES + 1)
    if not conteudo:
        raise ValueError("O arquivo enviado está vazio.")
    if len(conteudo) > MAX_UPLOAD_BYTES:
        raise ValueError("O arquivo excede o limite de 10 MB.")

    caminho_temporario: Path | None = None

    try:
        with tempfile.NamedTemporaryFile(delete=False, suffix=extensao) as temporario:
            temporario.write(conteudo)
            caminho_temporario = Path(temporario.name)

        texto = (
            _extrair_docx(caminho_temporario)
            if extensao == ".docx"
            else _extrair_doc(caminho_temporario)
        )

        if not texto:
            raise ValueError("O contrato não contém texto legível para análise.")

        return texto
    finally:
        if caminho_temporario:
            caminho_temporario.unlink(missing_ok=True)
