import io
import unittest
from unittest.mock import patch

from docx import Document
from fastapi.testclient import TestClient

from app.main import agente, app


class AgenteJuridicoAppTestCase(unittest.TestCase):
    def setUp(self):
        self.client = TestClient(app)

    @staticmethod
    def contrato_docx() -> io.BytesIO:
        arquivo = io.BytesIO()
        documento = Document()
        documento.add_heading("Contrato de prestação de serviços", level=1)
        documento.add_paragraph("Cláusula 1. A contratada prestará os serviços descritos.")
        documento.save(arquivo)
        arquivo.seek(0)
        return arquivo

    def test_pagina_inicial_e_estilos_respondem(self):
        resposta = self.client.get("/")
        self.assertEqual(resposta.status_code, 200)
        self.assertIn("Assistente de Revisão Contratual".encode("utf-8"), resposta.content)
        self.assertIn(b"/static/styles.css", resposta.content)

        estilos = self.client.get("/static/styles.css")
        self.assertEqual(estilos.status_code, 200)
        self.assertIn(b"--primary", estilos.content)

    def test_healthcheck(self):
        resposta = self.client.get("/health")
        self.assertEqual(resposta.status_code, 200)
        self.assertEqual(resposta.json(), {"status": "ok"})

    def test_rejeita_extensao_nao_permitida(self):
        resposta = self.client.post(
            "/analisar",
            files={"file": ("contrato.txt", b"conteudo", "text/plain")},
            data={"confirmacao": "confirmado"},
        )
        self.assertEqual(resposta.status_code, 400)
        self.assertIn("formatos .doc ou .docx".encode("utf-8"), resposta.content)

    def test_processa_docx_sem_chamar_servico_externo(self):
        with patch.object(
            agente,
            "analisar_bai",
            return_value="1. RISCOS JURÍDICOS\nNenhum risco crítico identificado.",
        ):
            resposta = self.client.post(
                "/analisar",
                files={
                    "file": (
                        "contrato.docx",
                        self.contrato_docx(),
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    )
                },
                data={"confirmacao": "confirmado"},
            )

        self.assertEqual(resposta.status_code, 200)
        self.assertIn("Relatório técnico".encode("utf-8"), resposta.content)
        self.assertIn("Nenhum risco crítico".encode("utf-8"), resposta.content)

    def test_escapa_html_retornado_pela_ia(self):
        with patch.object(
            agente,
            "analisar_bai",
            return_value="<script>alert('teste')</script>",
        ):
            resposta = self.client.post(
                "/analisar",
                files={
                    "file": (
                        "contrato.docx",
                        self.contrato_docx(),
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    )
                },
                data={"confirmacao": "confirmado"},
            )

        self.assertEqual(resposta.status_code, 200)
        self.assertIn(b"&lt;script&gt;alert", resposta.content)
        self.assertNotIn(b"<script>alert('teste')</script>", resposta.content)

    def test_nao_exige_confirmacao_para_analisar(self):
        with patch.object(
            agente,
            "analisar_bai",
            return_value="Relatório sem confirmação.",
        ):
            resposta = self.client.post(
                "/analisar",
                files={
                    "file": (
                        "contrato.docx",
                        self.contrato_docx(),
                        "application/vnd.openxmlformats-officedocument."
                        "wordprocessingml.document",
                    )
                },
            )

        self.assertEqual(resposta.status_code, 200)

    def test_respostas_possuem_cabecalhos_de_seguranca(self):
        resposta = self.client.get("/")
        self.assertEqual(resposta.headers["x-frame-options"], "DENY")
        self.assertEqual(resposta.headers["x-content-type-options"], "nosniff")
        self.assertEqual(resposta.headers["cache-control"], "no-store")
        self.assertIn("default-src 'self'", resposta.headers["content-security-policy"])


if __name__ == "__main__":
    unittest.main()
